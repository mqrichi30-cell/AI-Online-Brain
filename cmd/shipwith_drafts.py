#!/usr/bin/env python3
"""
AWG - Ship with POs Report  ·  one-command draft generator
==========================================================
Reads the "Consolidations" report (the .xlsx attached to the
"AWG - Ship with POs Repot" email) plus the contacts workbook, and writes one
ready-to-send Outlook draft (.eml) per recipient into an output folder.

Each .eml carries the `X-Unsent: 1` header, so double-clicking it in Windows
opens it in Outlook as an editable draft (review, then Send).

Rules
-----
- Signature: generic plain text (no personal signature).
- Matched Ship-to  -> To = the contact's email.
- Unmatched Ship-to -> To = pgcustservw2.im@pg.com, and a red
  "Please add contact for <name>." note is added at the top of the email.
- CC is fixed to the original distribution list.

Usage
-----
    python shipwith_drafts.py <report.xlsx> [contacts.xlsx] [-o OUTDIR]

Outputs (default OUTDIR = ./drafts_out):
    drafts_out/*.eml           one draft per recipient (open in Outlook)
    drafts_out/draft_emails.json
    drafts_out/preview_all.html
"""

import argparse
import json
import os
import re
from email.message import EmailMessage
from html import escape

import openpyxl

# --- Constants (mirror the original VBA) ---------------------------------------

SHARED_MAILBOX = "pgcustservw2.im@pg.com"
CC_LIST = "marquardt.jw@pg.com; jackson.vs@pg.com"

INTRO_HTML = (
    "Dear Team,<br><br>"
    "The following PO is undersized and requires a Ship With to make a full "
    "truckload. Do you prefer to write a Ship With PO we can consolidate OR do "
    "you prefer we add or increase product already on the PO?<br><br>"
)

# Generic signature (plain text look, no personal data).
SIGNATURE_HTML = (
    "<br>Regards,<br>"
    "pgcustservw2.im@pg.com<br>"
    "NA Order Management | Regional<br>"
)

# Column layout of the attachment "Consolidations" sheet (0-based).
COL_DB, COL_CREATED_ON, COL_GI, COL_RDD, COL_PLANT, COL_MIX, COL_SHIPTO = range(7)
COL_DOC, COL_PO, COL_SHIPCOND, COL_WEIGHT, COL_COF, COL_FP, COL_TM, COL_STATUS = range(7, 15)


def norm_name(value):
    if value is None:
        return ""
    return " ".join(str(value).split()).upper()


def fmt_date(value):
    if value is None:
        return ""
    try:
        return f"{value.month}/{value.day}/{value.year}"
    except AttributeError:
        return str(value)


def fmt_num(value, decimals=3):
    if value in (None, ""):
        return ""
    try:
        return f"{float(value):,.{decimals}f}"
    except (ValueError, TypeError):
        return str(value)


def is_data_row(row):
    return bool(str(row[COL_DB] or "").strip()) and bool(str(row[COL_SHIPTO] or "").strip())


def parse_consolidations(ws):
    blocks, current, number = [], [], 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        row = list(row) + [None] * (COL_STATUS + 1 - len(row))
        if is_data_row(row):
            current.append(row)
        elif current:
            number += 1
            blocks.append(_build_block(number, current))
            current = []
    if current:
        number += 1
        blocks.append(_build_block(number, current))
    return blocks


def _build_block(number, rows):
    return {
        "number": number,
        "ship_to": rows[0][COL_SHIPTO],
        "rows": [
            {
                "created_on": r[COL_CREATED_ON], "gi": r[COL_GI], "rdd": r[COL_RDD],
                "plant": r[COL_PLANT], "mix": r[COL_MIX], "ship_to": r[COL_SHIPTO],
                "doc": r[COL_DOC], "po": r[COL_PO], "ship_cond": r[COL_SHIPCOND],
                "weight": r[COL_WEIGHT], "fp": r[COL_FP], "status": r[COL_STATUS],
            }
            for r in rows
        ],
    }


def block_table_html(block):
    head = (
        "<tr>"
        "<th style='background-color:yellow;'>Created On</th>"
        "<th style='background-color:yellow;'>Goods Issue Date</th>"
        "<th style='background-color:yellow;'>Request Delivery Day</th>"
        "<th style='background-color:yellow;'>Plant</th>"
        "<th>Material Mix</th><th>Ship to Name</th><th>Document Number</th>"
        "<th>Purchase Order Number</th><th>Shipping Condition</th>"
        "<th style='background-color:yellow;'>Gross Weight</th>"
        "<th style='background-color:yellow;'>FP</th><th>Description</th>"
        "</tr>"
    )
    body, total_weight, total_fp = [], 0.0, 0.0
    for r in block["rows"]:
        for k, acc in (("weight", "w"), ("fp", "f")):
            try:
                v = float(r[k] or 0)
            except (ValueError, TypeError):
                v = 0.0
            if acc == "w":
                total_weight += v
            else:
                total_fp += v
        body.append(
            "<tr>"
            f"<td>{escape(fmt_date(r['created_on']))}</td>"
            f"<td>{escape(fmt_date(r['gi']))}</td>"
            f"<td>{escape(fmt_date(r['rdd']))}</td>"
            f"<td>{escape(str(r['plant'] or ''))}</td>"
            f"<td>{escape(str(r['mix'] or ''))}</td>"
            f"<td>{escape(str(r['ship_to'] or ''))}</td>"
            f"<td>{escape(str(r['doc'] or ''))}</td>"
            f"<td>{escape(str(r['po'] or ''))}</td>"
            f"<td>{escape(str(r['ship_cond'] or ''))}</td>"
            f"<td>{escape(fmt_num(r['weight']))}</td>"
            f"<td>{escape(fmt_num(r['fp']))}</td>"
            f"<td>{escape(str(r['status'] or ''))}</td>"
            "</tr>"
        )
    totals = (
        "<tr><td colspan='9'><strong>Totals:</strong></td>"
        f"<td><strong>{total_weight:,.2f}</strong></td>"
        f"<td><strong>{total_fp:,.2f}</strong></td><td></td></tr>"
    )
    return (
        f"<h3>Consolidation Number: {block['number']}</h3>"
        "<table border='1' cellpadding='5' cellspacing='0'>"
        + head + "".join(body) + totals + "</table><br>"
    )


def build_drafts(blocks, contacts):
    grouped, unmatched = {}, []
    for block in blocks:
        email = contacts.get(norm_name(block["ship_to"]), "")
        matched = bool(email)
        if matched:
            key = email.lower()
        else:
            unmatched.append(block["ship_to"])
            key = "__unmatched__" + norm_name(block["ship_to"])
            email = SHARED_MAILBOX          # route missing contacts to the shared box
        g = grouped.setdefault(key, {"email": email, "matched": matched, "names": [], "blocks": []})
        if block["ship_to"] not in g["names"]:
            g["names"].append(block["ship_to"])
        g["blocks"].append(block)

    drafts = []
    for g in grouped.values():
        note = ""
        if not g["matched"]:
            missing = ", ".join(g["names"])
            note = (
                "<p style='color:#b00020;font-weight:bold;'>"
                f"Please add contact for {escape(missing)}.</p>"
            )
        body = note + INTRO_HTML
        for block in g["blocks"]:
            body += block_table_html(block)
        body += SIGNATURE_HTML
        drafts.append(
            {
                "to": g["email"],
                "cc": CC_LIST,
                "subject": "SHIP WITH NEEDED - " + "/".join(g["names"]),
                "htmlBody": "<html><body>" + body + "</body></html>",
                "matched": g["matched"],
            }
        )
    return drafts, unmatched


def load_contacts(path):
    if not path or not os.path.exists(path):
        return {}
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip().lower() for c in ws[1]]
    name_idx = next((i for i, h in enumerate(headers) if "name" in h), None)
    email_idx = next((i for i, h in enumerate(headers) if "email" in h or "mail" in h), None)
    contacts = {}
    if name_idx is None or email_idx is None:
        return contacts
    for row in ws.iter_rows(min_row=2, values_only=True):
        if name_idx < len(row) and email_idx < len(row) and row[name_idx] and row[email_idx]:
            contacts[norm_name(row[name_idx])] = str(row[email_idx]).strip()
    return contacts


def safe_filename(text):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("_")[:120] or "draft"


def write_eml(draft, path):
    msg = EmailMessage()
    msg["Subject"] = draft["subject"]
    if draft["to"]:
        msg["To"] = draft["to"]
    if draft["cc"]:
        msg["Cc"] = draft["cc"]
    msg["From"] = SHARED_MAILBOX
    msg["X-Unsent"] = "1"          # makes Outlook open the .eml as an editable draft
    msg.set_content("This message requires an HTML-capable mail client.")
    msg.add_alternative(draft["htmlBody"], subtype="html")
    with open(path, "wb") as f:
        f.write(bytes(msg))


def write_preview(drafts, path):
    cards = []
    for m in drafts:
        to = m["to"] or "(blank)"
        tag = "" if m["matched"] else " <span style='color:#b00020'>· NO CONTACT</span>"
        cards.append(
            f"<div class='card'><div class='hdr'>"
            f"<div><span class='lbl'>To:</span> {escape(to)}{tag}</div>"
            f"<div><span class='lbl'>Cc:</span> {escape(m['cc'])}</div>"
            f"<div><span class='lbl'>Subject:</span> <b>{escape(m['subject'])}</b></div>"
            f"</div><div class='body'>{m['htmlBody']}</div></div>"
        )
    html = (
        "<div style=\"font-family:Segoe UI,Arial,sans-serif;max-width:1000px;margin:0 auto;padding:16px\">"
        f"<h1 style=\"font-size:20px\">Preview - Ship with POs drafts ({len(drafts)})</h1>"
        "<style>.card{border:1px solid #ddd;border-radius:8px;margin:16px 0;overflow:hidden}"
        ".hdr{background:#f5f6f8;padding:12px 16px;border-bottom:1px solid #e2e2e2;font-size:13px;line-height:1.6}"
        ".lbl{display:inline-block;min-width:56px;color:#666;font-weight:600}"
        ".body{padding:16px;font-size:13px;overflow-x:auto}"
        ".body table{border-collapse:collapse;font-size:12px}"
        ".body th,.body td{border:1px solid #999;padding:4px 8px;text-align:left}"
        ".body h3{font-size:14px;margin:14px 0 6px}</style>"
        + "".join(cards) + "</div>"
    )
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)


def main():
    ap = argparse.ArgumentParser(description="Generate Ship-with-POs Outlook drafts (.eml).")
    ap.add_argument("report", help="The Consolidations report .xlsx (email attachment).")
    ap.add_argument("contacts", nargs="?", help="Contacts workbook .xlsx (Ship-to name + email).")
    ap.add_argument("-o", "--out", default="drafts_out", help="Output folder (default: drafts_out).")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.report, data_only=True)
    ws = wb["Consolidations"] if "Consolidations" in wb.sheetnames else wb.active

    blocks = parse_consolidations(ws)
    contacts = load_contacts(args.contacts)
    drafts, unmatched = build_drafts(blocks, contacts)

    os.makedirs(args.out, exist_ok=True)
    for old in os.listdir(args.out):
        if old.endswith(".eml"):
            os.remove(os.path.join(args.out, old))

    for i, d in enumerate(drafts, 1):
        who = d["to"] if d["matched"] else "NO_CONTACT_" + "_".join(d["subject"].split(" - ")[1:])
        fname = f"{i:02d}_{safe_filename(who)}.eml"
        write_eml(d, os.path.join(args.out, fname))

    with open(os.path.join(args.out, "draft_emails.json"), "w", encoding="utf-8") as f:
        json.dump(drafts, f, indent=2, default=str)
    write_preview(drafts, os.path.join(args.out, "preview_all.html"))

    print(f"Parsed {len(blocks)} consolidations -> {len(drafts)} draft(s).")
    print(f"Contacts loaded: {len(contacts)}")
    if unmatched:
        print(f"No contact ({len(unmatched)}) -> routed to {SHARED_MAILBOX} with 'Please add contact' note:")
        for n in sorted(set(unmatched)):
            print(f"    - {n}")
    print(f"Output folder: {os.path.abspath(args.out)}")
    print("Open the .eml files in Outlook (they open as editable drafts), review, and Send.")


if __name__ == "__main__":
    main()
