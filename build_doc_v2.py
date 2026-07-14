# -*- coding: utf-8 -*-
"""Version 2 (sin referencias academicas externas): mismo contenido y formato APA 7,
citando solo las tres marcas referentes reales."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"
SIZE = 12

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(SIZE)
normal.font.color.rgb = RGBColor(0, 0, 0)
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
pf.space_before = Pt(0)
pf.space_after = Pt(0)

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def add_page_number(section):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    run.font.name = FONT
    run.font.size = Pt(SIZE)
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)
add_page_number(doc.sections[0])

def set_run(run, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(SIZE)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)

def title(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(text), bold=True); return p

def h1(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(text), bold=True); return p

def h2(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(p.add_run(text), bold=True); return p

def h3(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(p.add_run(text), bold=True, italic=True); return p

def para(runs, justify=True, first_line_indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    if first_line_indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    if isinstance(runs, str):
        runs = [(runs, False, False)]
    for text, bold, italic in runs:
        set_run(p.add_run(text), bold=bold, italic=italic)
    return p

def bullet(runs):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    if isinstance(runs, str):
        runs = [(runs, False, False)]
    for text, bold, italic in runs:
        set_run(p.add_run(text), bold=bold, italic=italic)
    return p

# ---------------- TITULO ----------------
title("Frooty: Propuesta de Identidad Visual para una Marca Ficticia de Bebidas "
      "Naturales Funcionales")
title("Avance 1 — Análisis, Investigación e Ideas Iniciales (Fase 1)")
doc.add_paragraph()

# ---------------- INTRODUCCION ----------------
h1("Introducción")
para("El presente documento corresponde a la primera fase del Proyecto Integrador del curso "
     "Introducción al Diseño Gráfico, cuyo propósito es desarrollar una propuesta de identidad "
     "visual (brand identity) para una marca ficticia, aplicando los fundamentos del diseño y de "
     "la comunicación de marca mediante la metodología de Aprendizaje Basado en Proyectos. La "
     "marca seleccionada por el equipo, denominada Frooty, se ubica dentro del ámbito comercial, "
     "específicamente en la categoría de bebidas saludables.")
para("Una marca es mucho más que un nombre o un logotipo: es la manera en que un producto se "
     "identifica y se distingue de los demás, y la identidad visual es la expresión tangible de "
     "esa marca, aquello que el público puede ver, reconocer y recordar. Bajo esta lógica, el "
     "presente avance reúne la investigación de referentes, la definición conceptual de la marca, "
     "el público meta, el moodboard inicial y las primeras propuestas de logotipo que servirán de "
     "base para la identidad visual final.")

# ---------------- 1. DESCRIPCION ----------------
h1("Descripción General de la Marca")
para([("Frooty es una marca costarricense ficticia dedicada a la creación de ", False, False),
      ("bebidas 100 % naturales", True, False),
      (". Se especializa en una línea premium de batidos (smoothies) funcionales y balanceados, "
       "acompañada de wellness shots prensados en frío —concentrados de jengibre, cúrcuma y limón "
       "orientados al fortalecimiento del sistema inmune—. Su propósito es facilitar el bienestar "
       "diario, transformando la nutrición sana y real en una experiencia genuinamente "
       "disfrutable. La marca no promueve dietas restrictivas ni sabores aburridos; su idea "
       "central sostiene que cuidar el cuerpo a través de ingredientes vivos, puros y locales "
       "puede ser un hábito delicioso, práctico y lleno de energía.", False, False)])
para("Esta idea central articula todo el proyecto: cada decisión de nombre, color, tipografía y "
     "símbolo buscará comunicar frescura, naturalidad, energía y disfrute, evitando la estética "
     "clínica o excesivamente sobria que suele asociarse a lo “saludable”.")

# ---------------- 2. REFERENTES ----------------
h1("Referentes Visuales")
para("La investigación de referentes permite identificar los códigos visuales propios de la "
     "categoría y detectar oportunidades de diferenciación. Para ello se analizaron tres marcas "
     "reales del mercado costarricense de bebidas y alimentos saludables, observando de forma "
     "directa su comunicación en sus canales oficiales.")

h2("Referente 1: Raw Co. CR")
para("Raw Co. CR es una marca costarricense que ofrece jugos prensados en frío, smoothies y "
     "otros productos naturales, promoviendo un estilo de vida basado en ingredientes frescos y "
     "nutritivos (Raw Co. CR, s. f.). En el plano visual destaca por una paleta de colores "
     "naturales —principalmente verde, blanco y tonos tierra—, un diseño minimalista que "
     "transmite limpieza y frescura, una tipografía moderna y fácil de leer, y el uso de imágenes "
     "que resaltan los ingredientes naturales.")
para("Esta identidad funciona porque comunica salud, frescura y bienestar de manera coherente: la "
     "sobriedad del diseño y el predominio del verde refuerzan la percepción de calidad y "
     "naturalidad, conectando con consumidores que buscan un estilo de vida saludable.")

h2("Referente 2: LATICA")
para("LATICA es una marca costarricense de bebidas elaboradas con frutas, enfocada en transmitir "
     "frescura, sabor y un estilo de vida saludable (LATICA, s. f.). Sus elementos visuales se "
     "apoyan en colores vibrantes inspirados en las frutas y la naturaleza, fotografías e "
     "ilustraciones de fruta que resaltan los ingredientes, una tipografía moderna y sencilla de "
     "identificar, y un diseño fresco y llamativo.")
para("Su identidad visual funciona porque los colores vivos y los elementos frutales hacen el "
     "producto atractivo y refuerzan la idea de una bebida saludable y divertida, convirtiéndola "
     "en una inspiración directa para Frooty.")

h2("Referente 3: Pura Frutika")
para("Pura Frutika es una marca costarricense que ofrece jugos, batidos, bowls y otros alimentos "
     "elaborados con frutas frescas, con una propuesta centrada en la alimentación saludable "
     "mediante productos naturales (Pura Frutika, s. f.). Visualmente emplea colores vibrantes "
     "inspirados en las frutas, fotografías de fruta fresca y de productos preparados, una "
     "tipografía moderna y legible, y un diseño limpio que transmite frescura y un estilo de vida "
     "saludable.")
para("Su identidad comunica frescura, naturalidad y bienestar, y al mostrar una variedad de "
     "productos saludables proyecta una marca cercana y atractiva para consumidores que buscan "
     "opciones nutritivas. La combinación de fotografía apetitosa con una gráfica alegre logra "
     "que lo saludable se perciba como algo deseable y no como una obligación.")

h2("Síntesis de los referentes")
para("Los tres referentes comparten un mismo lenguaje visual de categoría: predominio de colores "
     "asociados a la fruta y la naturaleza, fotografía protagonizada por los ingredientes, "
     "tipografías modernas y legibles, y una atmósfera de frescura. Frooty retoma estos códigos, "
     "pero busca diferenciarse a través de una personalidad más juvenil, colorida y emocionalmente "
     "cercana, alineada con la idea de un bienestar disfrutable.")

# ---------------- 3. DEFINICION MARCA ----------------
h1("Definición de la Marca")
para("La misión, la visión y los valores constituyen el núcleo estratégico sobre el cual se "
     "construye la identidad de la marca; de su claridad depende la coherencia de toda la "
     "comunicación visual posterior.")

h3("Misión")
para("Nutrir el bienestar integral de nuestros consumidores a través de combinaciones innovadoras "
     "de frutas, vegetales y superalimentos locales, elaborados bajo los más altos estándares de "
     "pureza y frescura, facilitando un estilo de vida saludable, activo y disfrutable en cada "
     "sorbo.")

h3("Visión")
para("Ser, para el año 2029, la marca líder y referente en el mercado nacional de bebidas "
     "funcionales y wellness, reconocida por transformar la percepción de la alimentación "
     "saludable mediante la innovación constante en sabores, la sostenibilidad de nuestros "
     "procesos y el impacto positivo en la salud de nuestra comunidad.")

h3("Valores")
bullet([("Pureza sin filtros: ", True, False),
        ("productos 100 % naturales. Lo que se ve en la etiqueta es exactamente lo que entra al "
         "cuerpo.", False, False)])
bullet([("Bienestar disfrutable (enjoyable wellness): ", True, False),
        ("equilibrio perfecto entre el máximo valor nutricional y un sabor extraordinario.",
         False, False)])
bullet([("Consciencia sostenible: ", True, False),
        ("apoyo activo a los agricultores locales para el abastecimiento y optimización de los "
         "procesos para reducir la huella ambiental.", False, False)])
bullet([("Innovación funcional: ", True, False),
        ("investigación y experimentación constantes en nuevas combinaciones y beneficios.",
         False, False)])
para("En conjunto, estos valores perfilan una personalidad de marca alegre, honesta y cercana: "
     "combina la sinceridad de lo natural con el entusiasmo de una propuesta disfrutable e "
     "innovadora. Esta dualidad será determinante para orientar las decisiones cromáticas y "
     "tipográficas del proyecto.")

# ---------------- 4. PUBLICO META ----------------
h1("Público Meta")
para("Frooty está dirigida principalmente a jóvenes y adultos de entre 18 y 30 años, en especial "
     "estudiantes universitarios y profesionales jóvenes de zonas urbanas de Costa Rica, con un "
     "estilo de vida activo y en búsqueda de alternativas saludables para su alimentación diaria.")
para("Este público se caracteriza por valorar el bienestar físico, el consumo de productos "
     "naturales y la practicidad en sus compras. Le interesa mantener hábitos saludables, realizar "
     "actividad física y elegir bebidas elaboradas con ingredientes frescos, sin colorantes ni "
     "conservantes artificiales. Además, utiliza con frecuencia redes sociales como Instagram y "
     "TikTok para descubrir nuevas marcas y conocer recomendaciones, por lo que la presencia "
     "digital y el atractivo visual del empaque resultan decisivos.")

tp = doc.add_paragraph(); set_run(tp.add_run("Tabla 1"), bold=True)
tp2 = doc.add_paragraph(); set_run(tp2.add_run("Perfil del consumidor meta de Frooty"), italic=True)

perfil = [
    ("Aspecto", "Descripción"),
    ("Edad", "18 a 30 años"),
    ("Género", "Hombres y mujeres"),
    ("Ubicación", "Principalmente zonas urbanas de Costa Rica"),
    ("Ocupación", "Estudiantes universitarios y profesionales jóvenes"),
    ("Intereses", "Alimentación saludable, ejercicio, bienestar, sostenibilidad, productos "
                  "naturales y redes sociales"),
    ("Hábitos de compra", "Prefieren bebidas naturales, revisan los ingredientes antes de "
                          "comprar y valoran la calidad, el diseño del empaque y las marcas con "
                          "propósito"),
]
table = doc.add_table(rows=len(perfil), cols=2)
table.style = "Table Grid"
for i, (a, b) in enumerate(perfil):
    cells = table.rows[i].cells
    for j, txt in enumerate((a, b)):
        cells[j].width = Inches(1.8 if j == 0 else 4.2)
        cp = cells[j].paragraphs[0]
        cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        set_run(cp.add_run(txt), bold=(i == 0))
note = doc.add_paragraph()
set_run(note.add_run("Nota. "), italic=True)
set_run(note.add_run("Elaboración propia del equipo a partir de la definición de la marca Frooty."))
doc.add_paragraph()
para("Este perfil orienta las decisiones de diseño: al tratarse de un público joven, digital y "
     "sensible a la estética, la marca debe verse fresca, colorida y coherente en cada punto de "
     "contacto, desde el empaque hasta las redes sociales.")

# ---------------- 5. MOODBOARD ----------------
h1("Moodboard Inicial")
para("El moodboard es un collage de imágenes, texturas, colores y tipografías que comunica, de un "
     "solo vistazo, la sensación y la dirección visual de un proyecto. El moodboard de Frooty "
     "reúne fotografías de frutas, bebidas frescas, escenas de verano y playa, y empaques de "
     "estética juvenil, con el fin de fijar el tono emocional de la marca antes de tomar "
     "decisiones definitivas.")

h2("Colores")
para("La paleta propuesta es tropical y vibrante: verdes frescos, amarillos, naranjas, un rosa/"
     "fucsia enérgico y toques de celeste. El verde aporta las asociaciones de naturalidad, "
     "frescura y salud propias de la categoría, mientras que los tonos cálidos y saturados "
     "(amarillo, naranja y fucsia) transmiten energía, alegría y apetito, reforzando el atributo "
     "de “bienestar disfrutable”.")

h2("Tipografía")
para("Se propone combinar una tipografía redondeada y amigable —para el nombre de la marca— con "
     "una tipografía de palo seco (sans serif) moderna y legible para los textos secundarios. Las "
     "formas redondeadas comunican cercanía, juventud y calidez, valores acordes con el público "
     "meta, mientras que la sans serif garantiza legibilidad en empaques y en formatos digitales "
     "como Instagram y TikTok, canales prioritarios para el segmento.")

h2("Estilo visual general")
para("El estilo se define por formas orgánicas, elementos gráficos espontáneos (garabatos, flores "
     "y ondas), fotografía de fruta real y una composición dinámica y desenfadada. Esta dirección "
     "conecta con los códigos frescos y coloridos observados en los referentes, pero con una "
     "expresión más lúdica y emocional que diferencia a Frooty. La coherencia entre colores, "
     "tipografías y estilo asegura que el moodboard funcione como guía visual del proyecto y "
     "mantenga la alineación con la idea central de la marca.")

# ---------------- 6. LOGOTIPO ----------------
h1("Propuestas de Logotipo")
para("A partir de la investigación y del moodboard, el equipo desarrolló una exploración de al "
     "menos seis bocetos de logotipo para el nombre Frooty, con el objetivo de probar distintas "
     "personalidades gráficas antes de seleccionar una dirección definitiva. Las propuestas "
     "exploran tres tratamientos tipográficos principales, presentados además sobre dos contextos "
     "cromáticos (fondos amarillo y verde) para evaluar su comportamiento:")
bullet([("Versión manuscrita (script): ", True, False),
        ("letra cursiva fluida que aporta un carácter cercano, artesanal y natural.", False, False)])
bullet([("Versión redondeada tipo “bubble”: ", True, False),
        ("letras gruesas y redondeadas con un pequeño rostro sonriente, que refuerza la idea de "
         "disfrute y bienestar.", False, False)])
bullet([("Versión geométrica con detalle frutal: ", True, False),
        ("tipografía compacta en la que las letras “O” incorporan un pequeño elemento de "
         "fruta/hoja y una sonrisa integrada en la base, uniendo el concepto de fruta y "
         "positividad.", False, False)])
para("La elección de trazos redondeados, la sonrisa y el guiño frutal responde directamente a la "
     "personalidad definida para la marca: un carácter alegre, juvenil y cercano, combinado con "
     "la honestidad de lo natural. Al mismo tiempo, la búsqueda de una tipografía moderna y "
     "legible retoma el aprendizaje obtenido de los referentes analizados. Estos bocetos "
     "constituyen el punto de partida que se refinará en la fase final del proyecto hasta "
     "consolidar el logotipo definitivo.")

# ---------------- CONCLUSION ----------------
h1("Conclusión")
para("Esta primera fase permitió comprender el problema de diseño y construir las bases "
     "conceptuales y visuales de Frooty. El análisis de tres referentes reales del mercado "
     "costarricense reveló los códigos visuales de la categoría; la definición de misión, visión "
     "y valores estableció la personalidad de la marca; la caracterización del público meta "
     "confirmó la pertinencia de la propuesta; y el moodboard junto con los bocetos de logotipo "
     "tradujeron todo lo anterior en las primeras decisiones gráficas. Con estos fundamentos, el "
     "equipo cuenta con una dirección clara y justificada para desarrollar, en la Fase 2, la "
     "identidad visual final de la marca.")

# ---------------- REFERENCIAS (solo marcas referentes) ----------------
doc.add_page_break()
h1("Referencias")
refs = [
    [("LATICA [@drinklatica]. (s. f.). ", False, False),
     ("Publicaciones", False, True),
     (" [Perfil de Instagram]. Instagram. Recuperado el 14 de julio de 2026, de "
      "https://www.instagram.com/drinklatica/", False, False)],
    [("Pura Frutika. (s. f.). ", False, False),
     ("Alimentos saludables", False, True),
     (". Recuperado el 14 de julio de 2026, de https://purafrutika.com/", False, False)],
    [("Raw Co. CR. (s. f.). ", False, False),
     ("Inicio", False, True),
     (". Recuperado el 14 de julio de 2026, de https://rawcocr.com/", False, False)],
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    for text, bold, italic in ref:
        set_run(p.add_run(text), bold=bold, italic=italic)

out = "Frooty_Avance1_Documento_SinReferencias.docx"
doc.save(out)
print("Documento generado:", out)
