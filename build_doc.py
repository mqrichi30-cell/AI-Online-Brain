# -*- coding: utf-8 -*-
"""Genera el documento escrito (Avance 1 - Frooty) con formato APA 7."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"
SIZE = 12

doc = Document()

# --- Estilo base: Times New Roman 12, doble espacio, sangría e interlineado ---
normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(SIZE)
normal.font.color.rgb = RGBColor(0, 0, 0)
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
pf.space_before = Pt(0)
pf.space_after = Pt(0)

# Márgenes de 1 pulgada (APA)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# --- Número de página en el encabezado, alineado a la derecha (APA) ---
def add_page_number(section):
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    run.font.name = FONT
    run.font.size = Pt(SIZE)
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

add_page_number(doc.sections[0])

# --- Helpers ---
def set_run(run, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(SIZE)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)

def title(text):
    """Título del trabajo: negrita, centrado (APA, primera página de texto)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, bold=True)
    return p

def h1(text):
    """Encabezado nivel 1: centrado, negrita."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, bold=True)
    return p

def h2(text):
    """Encabezado nivel 2: al margen izquierdo, negrita."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run(r, bold=True)
    return p

def h3(text):
    """Encabezado nivel 3: al margen izquierdo, negrita cursiva."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run(r, bold=True, italic=True)
    return p

def para(runs, justify=True, first_line_indent=True):
    """runs = lista de (texto, bold, italic) o string simple."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    if first_line_indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    if isinstance(runs, str):
        runs = [(runs, False, False)]
    for item in runs:
        text, bold, italic = item
        r = p.add_run(text)
        set_run(r, bold=bold, italic=italic)
    return p

def bullet(runs):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    if isinstance(runs, str):
        runs = [(runs, False, False)]
    for text, bold, italic in runs:
        r = p.add_run(text)
        set_run(r, bold=bold, italic=italic)
    return p

# ============================================================
# TÍTULO (se repite en la primera página de texto; la portada la hace otra persona)
# ============================================================
title("Frooty: Propuesta de Identidad Visual para una Marca Ficticia de Bebidas "
      "Naturales Funcionales")
title("Avance 1 — Análisis, Investigación e Ideas Iniciales (Fase 1)")
doc.add_paragraph()

# ============================================================
# INTRODUCCIÓN
# ============================================================
h1("Introducción")
para("El presente documento corresponde a la primera fase del Proyecto Integrador del curso "
     "Introducción al Diseño Gráfico, cuyo propósito es desarrollar una propuesta de identidad "
     "visual (brand identity) para una marca ficticia, aplicando los fundamentos del diseño y "
     "de la comunicación de marca mediante la metodología de Aprendizaje Basado en Proyectos. "
     "La marca seleccionada por el equipo, denominada Frooty, se ubica dentro del ámbito "
     "comercial, específicamente en la categoría de bebidas saludables.")
para([("Antes de diseñar es necesario comprender qué se va a construir. La American Marketing "
       "Association define una marca como un ", False, False),
      ("“nombre, término, diseño, símbolo o cualquier otra característica que identifica el bien "
       "o servicio de un vendedor como distinto de los de otros vendedores”", False, True),
      (" (American Marketing Association [AMA], s. f.). La identidad visual es la expresión "
       "tangible de esa marca: aquello que el público puede ver, tocar y recordar, y que traduce "
       "la estrategia en signos gráficos coherentes (Wheeler, 2018). Bajo esta lógica, el "
       "presente avance reúne la investigación de referentes, la definición conceptual de la "
       "marca, el público meta, el moodboard inicial y las primeras propuestas de logotipo que "
       "servirán de base para la identidad visual final.", False, False)])

# ============================================================
# 1. DESCRIPCIÓN GENERAL DE LA MARCA
# ============================================================
h1("Descripción General de la Marca")
para([("Frooty es una marca costarricense ficticia dedicada a la creación de ", False, False),
      ("bebidas 100 % naturales", True, False),
      (". Se especializa en una línea premium de batidos (smoothies) funcionales y balanceados, "
       "acompañada de wellness shots prensados en frío —concentrados de jengibre, cúrcuma y limón "
       "orientados al fortalecimiento del sistema inmune—. Su propósito es facilitar el bienestar "
       "diario, transformando la nutrición sana y real en una experiencia genuinamente "
       "disfrutable. La marca no promueve dietas restrictivas ni sabores aburridos; su idea "
       "central sostiene que cuidar el cuerpo a través de ingredientes vivos, puros y locales "
       "puede ser un hábito delicioso, práctico y lleno de energía.", False, False)])
para("Esta idea central articula todo el proyecto: cada decisión de nombre, color, tipografía y "
     "símbolo buscará comunicar frescura, naturalidad, energía y disfrute, evitando la estética "
     "clínica o excesivamente sobria que suele asociarse a lo “saludable”.")

# ============================================================
# 2. REFERENTES VISUALES
# ============================================================
h1("Referentes Visuales")
para("La investigación de referentes es una etapa clave del proceso de branding, pues permite "
     "identificar los códigos visuales propios de la categoría y detectar oportunidades de "
     "diferenciación (Wheeler, 2018). Para ello se analizaron tres marcas reales del mercado "
     "costarricense de bebidas y alimentos saludables, observando de forma directa su "
     "comunicación en sus canales oficiales.")

h2("Referente 1: Raw Co. CR")
para("Raw Co. CR es una marca costarricense pionera en jugos prensados en frío (cold pressed), "
     "smoothies y alimentos saludables, que promueve un estilo de vida basado en ingredientes "
     "frescos y nutritivos (Raw Co. CR, s. f.). En el plano visual destaca por una paleta de "
     "colores naturales —principalmente verde, blanco y tonos tierra—, un diseño minimalista que "
     "transmite limpieza y frescura, una tipografía moderna y fácil de leer, y el uso de "
     "imágenes que resaltan los ingredientes naturales.")
para("Esta identidad funciona porque comunica salud, frescura y bienestar de manera coherente: "
     "la sobriedad del diseño y el predominio del verde refuerzan la percepción de calidad y "
     "naturalidad, conectando con consumidores que buscan un estilo de vida saludable. El verde, "
     "por su asociación cultural con lo natural y lo fresco, resulta especialmente eficaz para "
     "esta categoría (Labrecque & Milne, 2012).")

h2("Referente 2: LATICA")
para("LATICA es una marca costarricense de bebidas elaboradas con frutas (agua gasificada con "
     "sabor natural y electrólitos) enfocada en transmitir frescura, sabor y un estilo de vida "
     "saludable (LATICA, s. f.). Sus elementos visuales se apoyan en colores vibrantes inspirados "
     "en las frutas y la naturaleza, fotografías e ilustraciones de fruta que resaltan los "
     "ingredientes, una tipografía moderna y sencilla de identificar, y un diseño fresco y "
     "llamativo.")
para("Su identidad visual funciona porque los colores vivos y los elementos frutales hacen el "
     "producto atractivo y refuerzan la idea de una bebida saludable y divertida. Desde la "
     "psicología del color aplicada al marketing, los tonos altamente saturados amplifican rasgos "
     "de personalidad de marca como el entusiasmo y la energía (Labrecque & Milne, 2012), lo cual "
     "explica el impacto de su comunicación y la convierte en una inspiración directa para Frooty.")

h2("Referente 3: Pura Frutika")
para("Pura Frutika es una marca costarricense que ofrece jugos, batidos, bowls y otros alimentos "
     "elaborados con frutas frescas, con una propuesta centrada en la alimentación saludable "
     "mediante productos naturales (Pura Frutika, s. f.). Visualmente emplea colores vibrantes "
     "inspirados en las frutas, fotografías de fruta fresca y de productos preparados, una "
     "tipografía moderna y legible, y un diseño limpio que transmite frescura y un estilo de vida "
     "saludable.")
para("Su identidad comunica frescura, naturalidad y bienestar, y al mostrar una variedad de "
     "productos saludables proyecta una marca cercana y atractiva para consumidores que buscan "
     "opciones nutritivas. La combinación de fotografía apetitosa con una gráfica alegre logra "
     "que lo saludable se perciba como algo deseable y no como una obligación.")

h2("Síntesis de los referentes")
para("Los tres referentes comparten un mismo lenguaje visual de categoría: predominio de colores "
     "asociados a la fruta y la naturaleza, fotografía protagonizada por los ingredientes, "
     "tipografías modernas y legibles, y una atmósfera de frescura. Frooty retoma estos códigos, "
     "pero busca diferenciarse a través de una personalidad más juvenil, colorida y "
     "emocionalmente cercana, alineada con la idea de un bienestar disfrutable.")

# ============================================================
# 3. DEFINICIÓN DE LA MARCA
# ============================================================
h1("Definición de la Marca")
para("La misión, la visión y los valores constituyen el núcleo estratégico sobre el cual se "
     "construye la identidad; de su claridad depende la coherencia de toda la comunicación visual "
     "posterior (Wheeler, 2018). Asimismo, estos elementos definen la personalidad de la marca, "
     "entendida como el conjunto de rasgos humanos asociados a ella (Aaker, 1997).")

h3("Misión")
para("Nutrir el bienestar integral de nuestros consumidores a través de combinaciones "
     "innovadoras de frutas, vegetales y superalimentos locales, elaborados bajo los más altos "
     "estándares de pureza y frescura, facilitando un estilo de vida saludable, activo y "
     "disfrutable en cada sorbo.", first_line_indent=True)

h3("Visión")
para("Ser, para el año 2029, la marca líder y referente en el mercado nacional de bebidas "
     "funcionales y wellness, reconocida por transformar la percepción de la alimentación "
     "saludable mediante la innovación constante en sabores, la sostenibilidad de nuestros "
     "procesos y el impacto positivo en la salud de nuestra comunidad.", first_line_indent=True)

h3("Valores")
bullet([("Pureza sin filtros: ", True, False),
        ("productos 100 % naturales. Lo que se ve en la etiqueta es exactamente lo que entra al "
         "cuerpo.", False, False)])
bullet([("Bienestar disfrutable (enjoyable wellness): ", True, False),
        ("equilibrio perfecto entre el máximo valor nutricional y un sabor extraordinario.",
         False, False)])
bullet([("Consciencia sostenible: ", True, False),
        ("apoyo activo a los agricultores locales para el abastecimiento y optimización de los "
         "procesos para reducir la huella ambiental.", False, False)])
bullet([("Innovación funcional: ", True, False),
        ("investigación y experimentación constantes en nuevas combinaciones y beneficios.",
         False, False)])
para("En conjunto, estos valores perfilan una personalidad de marca que combina la sinceridad "
     "—asociada a la honestidad y la naturalidad de “pureza sin filtros”— con el entusiasmo "
     "—vinculado a lo disfrutable, innovador y enérgico— dos de las dimensiones descritas por "
     "Aaker (1997). Esta dualidad será determinante para orientar las decisiones cromáticas y "
     "tipográficas del proyecto.")

# ============================================================
# 4. PÚBLICO META
# ============================================================
h1("Público Meta")
para("La definición del público meta consiste en seleccionar el segmento de consumidores al que "
     "la marca dirigirá su oferta, a partir de criterios demográficos y psicográficos (Kotler & "
     "Armstrong, 2017). Frooty está dirigida principalmente a jóvenes y adultos de entre 18 y 30 "
     "años, en especial estudiantes universitarios y profesionales jóvenes de zonas urbanas de "
     "Costa Rica, con un estilo de vida activo y en búsqueda de alternativas saludables para su "
     "alimentación diaria.")
para("Este público se caracteriza por valorar el bienestar físico, el consumo de productos "
     "naturales y la practicidad en sus compras. Le interesa mantener hábitos saludables, "
     "realizar actividad física y elegir bebidas elaboradas con ingredientes frescos, sin "
     "colorantes ni conservantes artificiales. Además, utiliza con frecuencia redes sociales como "
     "Instagram y TikTok para descubrir nuevas marcas y conocer recomendaciones, por lo que la "
     "presencia digital y el atractivo visual del empaque resultan decisivos.")

# Tabla 1
tp = doc.add_paragraph()
r = tp.add_run("Tabla 1")
set_run(r, bold=True)
tp2 = doc.add_paragraph()
r2 = tp2.add_run("Perfil del consumidor meta de Frooty")
set_run(r2, italic=True)

perfil = [
    ("Aspecto", "Descripción"),
    ("Edad", "18 a 30 años"),
    ("Género", "Hombres y mujeres"),
    ("Ubicación", "Principalmente zonas urbanas de Costa Rica"),
    ("Ocupación", "Estudiantes universitarios y profesionales jóvenes"),
    ("Intereses", "Alimentación saludable, ejercicio, bienestar, sostenibilidad, productos "
                  "naturales y redes sociales"),
    ("Hábitos de compra", "Prefieren bebidas naturales, revisan los ingredientes antes de "
                          "comprar y valoran la calidad, el diseño del empaque y las marcas con "
                          "propósito"),
]
table = doc.add_table(rows=len(perfil), cols=2)
table.style = "Table Grid"
table.columns[0].width = Inches(1.8)
table.columns[1].width = Inches(4.2)
for i, (a, b) in enumerate(perfil):
    cells = table.rows[i].cells
    for j, txt in enumerate((a, b)):
        cells[j].width = Inches(1.8 if j == 0 else 4.2)
        cp = cells[j].paragraphs[0]
        cp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        rr = cp.add_run(txt)
        set_run(rr, bold=(i == 0))
# nota de tabla
note = doc.add_paragraph()
rn = note.add_run("Nota. ")
set_run(rn, italic=True)
rn2 = note.add_run("Elaboración propia del equipo a partir de la definición de la marca Frooty.")
set_run(rn2)
doc.add_paragraph()

h2("Contexto de mercado y oportunidad")
para("La elección de este público responde a una necesidad real observable en el país. El estudio "
     "de Guevara-Villalobos et al. (2019), realizado con 798 participantes del área urbana de "
     "Costa Rica entre 15 y 65 años en el marco del Estudio Latinoamericano de Nutrición y Salud "
     "(ELANS), concluyó que los hábitos alimentarios de la población urbana costarricense son "
     "poco variados, con un alto consumo de café, panes, arroz blanco y bebidas azucaradas, y un "
     "consumo insuficiente de frutas, vegetales y leguminosas. Este panorama evidencia una "
     "oportunidad clara para una marca como Frooty, que ofrece bebidas a base de fruta natural "
     "como alternativa a las bebidas azucaradas de consumo habitual.")
para("Esta oportunidad se refuerza con la tendencia de la “indulgencia saludable”, en la que el "
     "consumidor busca productos que sean, a la vez, saludables y placenteros, impulsando la "
     "demanda de alimentos y bebidas naturales y funcionales (Promotora del Comercio Exterior de "
     "Costa Rica [PROCOMER], s. f.). Frooty se inserta precisamente en esa intersección entre lo "
     "nutritivo y lo disfrutable, lo que hace coherente su propuesta con el comportamiento actual "
     "del mercado.")

# ============================================================
# 5. MOODBOARD INICIAL
# ============================================================
h1("Moodboard Inicial")
para("Un moodboard es un collage de imágenes, texturas, colores y tipografías que comunica, de un "
     "solo vistazo, la sensación y la dirección visual de un proyecto; es una herramienta que "
     "traduce en imágenes conceptos y emociones difíciles de expresar solo con palabras "
     "(Interaction Design Foundation [IDF], s. f.). El moodboard de Frooty reúne fotografías de "
     "frutas, bebidas frescas, escenas de verano y playa, y empaques de estética juvenil, con el "
     "fin de fijar el tono emocional de la marca antes de tomar decisiones definitivas.")

h2("Colores")
para("La paleta propuesta es tropical y vibrante: verdes frescos, amarillos, naranjas, un rosa/"
     "fucsia enérgico y toques de celeste. Esta selección se fundamenta en la psicología del "
     "color aplicada al marketing: el color influye de manera directa en la percepción de la "
     "personalidad de la marca, y la saturación y el brillo amplifican rasgos como el entusiasmo "
     "y la energía (Labrecque & Milne, 2012). El verde aporta las asociaciones de naturalidad, "
     "frescura y salud propias de la categoría, mientras que los tonos cálidos y saturados "
     "(amarillo, naranja y fucsia) transmiten energía, alegría y apetito, reforzando el atributo "
     "de “bienestar disfrutable”.")

h2("Tipografía")
para("Se propone combinar una tipografía redondeada y amigable —para el nombre de la marca— con "
     "una tipografía de palo seco (sans serif) moderna y legible para los textos secundarios. "
     "Las formas redondeadas comunican cercanía, juventud y calidez, valores acordes con el "
     "público meta, mientras que la sans serif garantiza legibilidad en empaques y en formatos "
     "digitales como Instagram y TikTok, canales prioritarios para el segmento.")

h2("Estilo visual general")
para("El estilo se define por formas orgánicas, elementos gráficos espontáneos (garabatos, "
     "flores y ondas), fotografía de fruta real y una composición dinámica y desenfadada. Esta "
     "dirección conecta con los códigos frescos y coloridos observados en los referentes, pero "
     "con una expresión más lúdica y emocional que diferencia a Frooty. La coherencia entre "
     "colores, tipografías y estilo asegura que el moodboard funcione como guía visual del "
     "proyecto y mantenga la alineación con la idea central de la marca (IDF, s. f.).")

# ============================================================
# 6. PROPUESTAS DE LOGOTIPO
# ============================================================
h1("Propuestas de Logotipo")
para("A partir de la investigación y del moodboard, el equipo desarrolló una exploración de al "
     "menos seis bocetos de logotipo para el nombre Frooty, con el objetivo de probar distintas "
     "personalidades gráficas antes de seleccionar una dirección definitiva. Las propuestas "
     "exploran tres tratamientos tipográficos principales, presentados además sobre dos contextos "
     "cromáticos (fondos amarillo y verde) para evaluar su comportamiento:")
bullet([("Versión manuscrita (script): ", True, False),
        ("letra cursiva fluida que aporta un carácter cercano, artesanal y natural.", False, False)])
bullet([("Versión redondeada tipo “bubble”: ", True, False),
        ("letras gruesas y redondeadas con un pequeño rostro sonriente, que refuerza la idea de "
         "disfrute y bienestar.", False, False)])
bullet([("Versión geométrica con detalle frutal: ", True, False),
        ("tipografía compacta en la que las letras “O” incorporan un pequeño elemento de "
         "fruta/hoja y una sonrisa integrada en la base, uniendo el concepto de fruta y "
         "positividad.", False, False)])
para("La elección de trazos redondeados, la sonrisa y el guiño frutal responde directamente a la "
     "personalidad definida para la marca: un carácter alegre, juvenil y cercano (entusiasmo) "
     "combinado con la honestidad de lo natural (sinceridad), en línea con las dimensiones de "
     "personalidad de marca de Aaker (1997). Al mismo tiempo, la búsqueda de una tipografía "
     "moderna y legible retoma el aprendizaje obtenido de los referentes analizados. Estos "
     "bocetos constituyen el punto de partida que se refinará en la fase final del proyecto "
     "hasta consolidar el logotipo definitivo.")

# ============================================================
# CONCLUSIÓN
# ============================================================
h1("Conclusión")
para("Esta primera fase permitió comprender el problema de diseño y construir las bases "
     "conceptuales y visuales de Frooty. El análisis de tres referentes reales del mercado "
     "costarricense reveló los códigos visuales de la categoría; la definición de misión, visión "
     "y valores estableció la personalidad de la marca; la caracterización del público meta, "
     "respaldada por datos de consumo nacionales, confirmó la pertinencia de la propuesta; y el "
     "moodboard junto con los bocetos de logotipo tradujeron todo lo anterior en las primeras "
     "decisiones gráficas. Con estos fundamentos, el equipo cuenta con una dirección clara y "
     "justificada para desarrollar, en la Fase 2, la identidad visual final de la marca.")

# ============================================================
# REFERENCIAS
# ============================================================
doc.add_page_break()
h1("Referencias")

refs = [
    [("Aaker, J. L. (1997). Dimensions of brand personality. ", False, False),
     ("Journal of Marketing Research, 34", False, True),
     ("(3), 347–356. https://doi.org/10.1177/002224379703400304", False, False)],

    [("American Marketing Association. (s. f.). ", False, False),
     ("Branding", False, True),
     (". Recuperado el 14 de julio de 2026, de https://www.ama.org/topics/brand-and-branding/",
      False, False)],

    [("Guevara-Villalobos, D., Céspedes-Vindas, C., Flores-Soto, N., Úbeda-Carrasquilla, L., "
      "Chinnock, A., & Gómez, G. (2019). Hábitos alimentarios de la población urbana "
      "costarricense. ", False, False),
     ("Acta Médica Costarricense, 61", False, True),
     ("(4), 152–159. https://www.scielo.sa.cr/scielo.php?script=sci_arttext&pid="
      "S0001-60022019000400152", False, False)],

    [("Interaction Design Foundation. (s. f.). ", False, False),
     ("Mood boards", False, True),
     (". Recuperado el 14 de julio de 2026, de "
      "https://www.interaction-design.org/literature/topics/mood-boards", False, False)],

    [("Kotler, P., & Armstrong, G. (2017). ", False, False),
     ("Fundamentos de marketing", False, True),
     (" (13.ª ed.). Pearson Educación.", False, False)],

    [("Labrecque, L. I., & Milne, G. R. (2012). Exciting red and competent blue: The importance "
      "of color in marketing. ", False, False),
     ("Journal of the Academy of Marketing Science, 40", False, True),
     ("(5), 711–727. https://doi.org/10.1007/s11747-010-0245-y", False, False)],

    [("LATICA [@drinklatica]. (s. f.). ", False, False),
     ("Publicaciones", False, True),
     (" [Perfil de Instagram]. Instagram. Recuperado el 14 de julio de 2026, de "
      "https://www.instagram.com/drinklatica/", False, False)],

    [("Promotora del Comercio Exterior de Costa Rica. (s. f.). ", False, False),
     ("El auge de la indulgencia saludable para el consumidor post-pandemia", False, True),
     (". Recuperado el 14 de julio de 2026, de https://www.procomer.com/noticia/"
      "el-auge-de-la-indulgencia-saludable-para-el-consumidor-post-pandemia/", False, False)],

    [("Pura Frutika. (s. f.). ", False, False),
     ("Alimentos saludables", False, True),
     (". Recuperado el 14 de julio de 2026, de https://purafrutika.com/", False, False)],

    [("Raw Co. CR. (s. f.). ", False, False),
     ("Inicio", False, True),
     (". Recuperado el 14 de julio de 2026, de https://rawcocr.com/", False, False)],

    [("Wheeler, A. (2018). ", False, False),
     ("Designing brand identity: An essential guide for the whole branding team", False, True),
     (" (5.ª ed.). John Wiley & Sons.", False, False)],
]

for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)  # sangría francesa
    for text, bold, italic in ref:
        r = p.add_run(text)
        set_run(r, bold=bold, italic=italic)

out = "Frooty_Avance1_Documento_APA7.docx"
doc.save(out)
print("Documento generado:", out)
